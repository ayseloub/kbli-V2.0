from pathlib import Path

import pytest

from kbli.documents import (DocumentError, Page, chunk_document, clean_pages, extract, fts_query,
                            structure)
from tests.make_fixtures import make_pdf

FIX = Path(__file__).parent / "fixtures" / "pp_tiruan.pdf"


@pytest.fixture(scope="module")
def pdf_pages():
    if not FIX.exists():
        make_pdf(FIX)
    return clean_pages(extract(FIX))


def test_pdf_header_footer_and_page_numbers_removed(pdf_pages):
    text = "\n".join(p.text for p in pdf_pages)
    assert "PRESIDEN REPUBLIK INDONESIA" not in text
    assert "SK No" not in text
    assert "- 2 -" not in text
    assert "Pasal 12" in text


def test_structure_tracks_bab_bagian_pasal_across_pages(pdf_pages):
    paras = structure(pdf_pages)
    p12 = next(p for p in paras if "huruf a berupa NIB" in p.text)
    assert p12.path == ("BAB II – PERIZINAN BERUSAHA BERBASIS RISIKO", "Bagian Kesatu – Umum", "Pasal 12")
    # baris PDF yang terpotong di tengah kalimat digabung
    assert any("untuk memulai dan menjalankan usaha" in p.text for p in paras)


def test_chunks_never_cross_bab_and_respect_size(pdf_pages):
    chunks = chunk_document(pdf_pages, "PP tiruan", max_words=60)
    assert len(chunks) >= 3
    for c in chunks:
        assert c.word_count <= 60 + 45          # + overlap & judul pasal
        assert c.embed_text.startswith("PP tiruan")
    babs = [c.heading_path[0] for c in chunks]
    assert babs == sorted(babs, key=babs.index)  # urutan BAB tidak bolak-balik
    sanksi = [c for c in chunks if "SANKSI" in c.heading_path[0]]
    assert sanksi and "Pasal 30" in sanksi[0].content and "NIB" not in sanksi[0].content


def test_overlap_keeps_pasal_label(pdf_pages):
    chunks = chunk_document(pdf_pages, "PP", max_words=60)
    with_overlap = [c for c in chunks if c.metadata["has_overlap"]]
    assert with_overlap
    for c in with_overlap:
        first, second = c.content.split("\n")[:2]
        assert first.startswith("Pasal") and second.startswith("…")


def test_docx_headings_and_tables(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Syarat Pendaftaran", level=1)
    d.add_paragraph("Pelaku usaha mikro cukup menyiapkan KTP dan alamat usaha.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Risiko", "Perizinan"
    t.cell(1, 0).text, t.cell(1, 1).text = "Rendah", "NIB"
    f = tmp_path / "sop.docx"
    d.save(f)
    chunks = chunk_document(clean_pages(extract(f)), "SOP")
    assert chunks[0].heading_path == ["Syarat Pendaftaran"]
    assert "Rendah | NIB" in chunks[0].content


def test_markdown_headings(tmp_path):
    f = tmp_path / "faq.md"
    f.write_text("# FAQ OSS\n\n## Apa itu NIB?\nNIB adalah identitas pelaku usaha.\n\n## Berapa biayanya?\nGratis.\n")
    chunks = chunk_document(clean_pages(extract(f)), "FAQ", min_words=1)
    paths = [c.heading_path for c in chunks]
    assert ["FAQ OSS", "Apa itu NIB?"] in paths or any("Apa itu NIB?" in " ".join(p) for p in paths)


def test_scanned_pdf_is_rejected_with_clear_message(tmp_path):
    from reportlab.pdfgen import canvas

    f = tmp_path / "scan.pdf"
    c = canvas.Canvas(str(f))
    c.rect(50, 50, 400, 600, fill=1)
    c.showPage()
    c.save()
    with pytest.raises(DocumentError, match="scan"):
        extract(f)


def test_unsupported_format(tmp_path):
    f = tmp_path / "x.xlsx"
    f.write_bytes(b"PK")
    with pytest.raises(DocumentError, match="belum didukung"):
        extract(f)


def test_long_paragraph_is_split():
    pages = [Page(1, "Pasal 1\n" + " ".join(f"kata{i}." for i in range(700)))]
    chunks = chunk_document(pages, "X", max_words=100)
    assert len(chunks) >= 7 and all(c.word_count <= 150 for c in chunks)


def test_fts_query():
    assert fts_query("gimana syarat perizinan usaha risiko rendah?") == "syarat:* | perizinan:* | usaha:* | risiko:* | rendah:*"
    assert fts_query("apa ya") is None
